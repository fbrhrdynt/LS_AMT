import io
from pathlib import Path
import re
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
from xml.sax.saxutils import escape

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Response
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image as RLImage,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from auth import get_current_user
from branding import get_pdf_brand_logo_bytes
from core import db
from license_service import require_feature_enabled

router = APIRouter(prefix="/api")

ASSET_DIR = Path(__file__).resolve().parent / "assets"
AMT_MARK_TAGLINE = ASSET_DIR / "amt-mark-tagline.png"


def _status_label(value):
    if value in ("Operational", "Green Tag / Ready"):
        return "Green Tag / Ready"
    if value in ("Under Maintenance", "Yellow Tag / Under Maintenance", "Red Tag / Under Maintenance"):
        return "Yellow Tag / Under Maintenance"
    if value in ("Out of Service", "Red Tag / Out of Service"):
        return "Red Tag / Out of Service"
    return value or ""


def _safe_excel(value):
    if value is None:
        return ""
    if isinstance(value, (int, float, bool)):
        return value
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _safe_pdf(value):
    return escape(str(value if value is not None else ""))


def _rx(text):
    return {"$regex": re.escape(text), "$options": "i"}


async def _timezone_name():
    settings = await db.settings.find_one(
        {"_id": "app"}, {"_id": 0, "timezone": 1}
    ) or {}
    return settings.get("timezone", "Asia/Jakarta")


def _local_now(timezone_name):
    try:
        tz = ZoneInfo(timezone_name)
    except Exception:
        tz = ZoneInfo("UTC")
    return datetime.now(timezone.utc).astimezone(tz)


async def _equipment_rows(params):
    query = {}
    q = params.get("q", "")
    if q:
        query["$or"] = [
            {"sap_no": _rx(q)}, {"mfg_no": _rx(q)}, {"name": _rx(q)},
            {"category": _rx(q)}, {"manufacturer": _rx(q)},
        ]
    if params.get("placement"):
        query["placement"] = params["placement"]
    if params.get("status"):
        query["operational_status"] = params["status"]

    records = await db.equipment.find(query, {"_id": 0}).sort("sap_no", 1).to_list(100000)
    job_ids = list({row.get("current_job_id") for row in records if row.get("current_job_id")})
    jobs = await db.jobs.find({"id": {"$in": job_ids}}, {"_id": 0}).to_list(10000)
    job_map = {j["id"]: j for j in jobs}

    rows = []
    for eq in records:
        placement = eq.get("placement") or "Base"
        detail = (eq.get("placement_detail") or "").strip()
        if placement == "Job":
            job = job_map.get(eq.get("current_job_id")) or {}
            location = " - ".join(
                p for p in ["Job", job.get("client_name") or "", job.get("site_location") or ""] if p
            ) or "Job"
        elif not detail or detail.lower() == placement.lower():
            location = placement
        else:
            location = f"{placement} - {detail}"
        rows.append([
            eq.get("sap_no"), eq.get("mfg_no"), eq.get("name"), eq.get("category"),
            eq.get("manufacturer"), eq.get("physical_condition"), location,
            _status_label(eq.get("operational_status")), eq.get("date_of_purchase"),
        ])

    return [
        "Asset / SAP No.", "Serial / Mfg No.", "Equipment", "Category",
        "Manufacturer", "Equipment Condition", "Current Location",
        "Operational Status", "Date of Purchase",
    ], rows


async def _maintenance_rows(params):
    query = {}

    maintenance_ids = [
        value.strip()
        for value in str(
            params.get("maintenance_ids", "")
            or ""
        ).split(",")
        if value.strip()
    ]

    if len(maintenance_ids) > 500:
        raise HTTPException(
            status_code=400,
            detail=(
                "A maximum of 500 maintenance "
                "records can be exported at once"
            ),
        )

    if maintenance_ids:
        query["id"] = {
            "$in": maintenance_ids
        }

    for key in ("status", "client_id", "job_id"):
        if params.get(key):
            query[key] = params[key]
    if params.get("sap_no"):
        query["sap_no"] = {"$regex": params["sap_no"], "$options": "i"}
    if params.get("type"):
        query["type_of_maintenance"] = {"$regex": params["type"], "$options": "i"}
    if params.get("failure"):
        query["failure_found"] = {"$regex": params["failure"], "$options": "i"}
    if params.get("technician"):
        tech = params["technician"]
        query["$or"] = [
            {"lead_technician": {"$regex": tech, "$options": "i"}},
            {"support_technicians": {"$regex": tech, "$options": "i"}},
        ]
    if params.get("date_from") or params.get("date_to"):
        dq = {}
        if params.get("date_from"):
            dq["$gte"] = params["date_from"]
        if params.get("date_to"):
            dq["$lte"] = params["date_to"]
        query["maintenance_date"] = dq

    records = await db.maintenance.find(query, {"_id": 0}).sort("maintenance_date", -1).to_list(100000)

    if params.get("serial_no"):
        eqs = await db.equipment.find(
            {"mfg_no": {"$regex": params["serial_no"], "$options": "i"}}, {"_id": 0, "id": 1}
        ).to_list(10000)
        ids = {e["id"] for e in eqs}
        records = [m for m in records if m.get("equipment_id") in ids]

    rows = []
    for m in records:
        purpose = m.get("maintenance_purpose") or ""
        field_name = m.get("field_name") or ""
        recorded_total = round(
            sum(
                float(part.get("cost") or 0)
                for part in (m.get("parts_consumed") or [])
                if (part.get("supply_source") or "Ex-Stock") in ("Purchase", "Warehouse")
            ),
            2,
        )
        rows.append([
            m.get("mnt_no"), m.get("maintenance_date"), m.get("date_closed"),
            m.get("sap_no"), m.get("equipment_name"), m.get("type_of_maintenance"),
            m.get("maintenance_category"), purpose, m.get("client_name"), field_name,
            m.get("lead_technician"), m.get("problem_damage"), m.get("failure_found"),
            m.get("status"), recorded_total,
        ])

    return [
        "Maintenance No.", "Maintenance Date", "Date Closed", "Asset / SAP No.",
        "Equipment", "Type", "Category", "Maintenance Purpose", "Client", "Field Name",
        "Lead Technician", "Observed Symptoms", "Failure Found", "Status",
        "Recorded Value Total",
    ], rows


def _calibration_expiry_status(
    expired_date,
):
    value = str(
        expired_date or ""
    ).strip()

    if not value:
        return "No Expiry"

    try:
        end = datetime.strptime(
            value[:10],
            "%Y-%m-%d",
        ).date()
    except ValueError:
        return "No Expiry"

    days = (
        end
        - datetime.now(
            timezone.utc
        ).date()
    ).days

    if days < 0:
        return "Expired"
    if days <= 30:
        return "Due Soon"
    return "Valid"


async def _calibration_rows(params):
    query = {
        "is_deleted": {
            "$ne": True
        }
    }

    q = params.get("q", "")
    if q:
        query["$or"] = [
            {"tool_id": _rx(q)},
            {"tool_name": _rx(q)},
            {"category": _rx(q)},
            {"equipment_sap_no": _rx(q)},
            {"equipment_name": _rx(q)},
            {"cert_number": _rx(q)},
            {"calibrated_by": _rx(q)},
        ]

    if params.get("category"):
        query["category"] = (
            params["category"]
        )

    records = await db.calibration_tools.find(
        query,
        {"_id": 0},
    ).sort(
        [
            ("category", 1),
            ("tool_id", 1),
        ]
    ).to_list(100000)

    rows = []
    wanted_status = params.get(
        "status"
    )

    for item in records:
        status = (
            _calibration_expiry_status(
                item.get(
                    "expired_date"
                )
            )
        )

        if (
            wanted_status
            and status
            != wanted_status
        ):
            continue

        frequency = ""
        if item.get(
            "frequency_value"
        ):
            unit = (
                "month(s)"
                if item.get(
                    "frequency_unit"
                )
                == "month"
                else "week(s)"
            )
            frequency = (
                f"{item.get('frequency_value')} "
                f"{unit}"
            )

        rows.append([
            item.get("tool_id"),
            item.get("tool_name"),
            item.get("category"),
            item.get("manufacturer"),
            item.get("model"),
            item.get("range_spec"),
            item.get("equipment_sap_no"),
            item.get("equipment_name"),
            item.get("calibration_date"),
            frequency,
            item.get("expired_date"),
            status,
            item.get("cert_number"),
            item.get("calibrated_by"),
            item.get("comments"),
        ])

    return [
        "Tool ID / Serial",
        "Tool",
        "Category",
        "Manufacturer",
        "Model",
        "Range / Set Pressure",
        "Assigned SAP",
        "Equipment",
        "Calibration Date",
        "Frequency",
        "Expired Date",
        "Status",
        "Certificate No.",
        "Calibrated By",
        "Comments / Notes",
    ], rows


async def _inventory_rows(params):
    query = {}
    q = params.get("q", "")
    if q:
        query["$or"] = [
            {"item_code": _rx(q)}, {"item_name": _rx(q)},
            {"part_number": _rx(q)}, {"storage_location": _rx(q)},
        ]
    if params.get("type"):
        query["type"] = params["type"]
    if params.get("low") in ("1", "true", "True"):
        query["$expr"] = {"$lte": ["$stock", "$min_stock"]}
    records = await db.inventory_items.find(query, {"_id": 0}).sort("item_code", 1).to_list(100000)
    return [
        "Item Code", "Item Name", "Type", "Part Number", "Unit", "Stock",
        "Minimum Stock", "Storage Location", "Unit Price",
    ], [[
        r.get("item_code"), r.get("item_name"), r.get("type"), r.get("part_number"),
        r.get("unit"), r.get("stock"), r.get("min_stock"), r.get("storage_location"),
        r.get("unit_price"),
    ] for r in records]


async def _client_rows(params):
    records = await db.clients.find({}, {"_id": 0}).sort("name", 1).to_list(100000)
    rows = []
    for c in records:
        count = await db.jobs.count_documents({"client_id": c["id"]})
        rows.append([c.get("name"), c.get("code"), c.get("contact"), count, c.get("notes")])
    return ["Client", "Code", "Contact", "Job Count", "Notes"], rows


async def _job_rows(params):
    records = await db.jobs.find({}, {"_id": 0}).sort("created_at", -1).to_list(100000)
    return [
        "Job ID", "Client", "Site", "Field Name", "Start Date", "End Date", "Status", "Notes",
    ], [[
        r.get("job_number"), r.get("client_name"), r.get("site_location"),
        r.get("field_name") or r.get("job_name"), r.get("start_date"), r.get("end_date"),
        r.get("status"), r.get("notes"),
    ] for r in records]


async def _audit_rows(params):
    query = {}
    if params.get("entity_type"):
        query["entity_type"] = params["entity_type"]
    records = await db.audit_logs.find(query, {"_id": 0}).sort("timestamp", -1).to_list(100000)
    return ["Timestamp", "Action", "Entity", "Entity ID", "Details", "User"], [[
        r.get("timestamp"), r.get("action"), r.get("entity_type"), r.get("entity_id"),
        r.get("details"), r.get("user_name"),
    ] for r in records]


async def _user_rows(params, user):
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Administrator access required")
    records = await db.users.find({}, {
        "_id": 0, "name": 1, "email": 1, "auth_provider": 1, "role": 1, "created_at": 1,
    }).sort("created_at", -1).to_list(100000)
    return ["Name", "Email", "Provider", "Role", "Created"], [[
        r.get("name"), r.get("email"), r.get("auth_provider"), r.get("role"), r.get("created_at"),
    ] for r in records]


async def _dashboard_rows(params):
    total = await db.equipment.count_documents({})
    green = await db.equipment.count_documents({"operational_status": "Operational"})
    red = await db.equipment.count_documents({"operational_status": "Under Maintenance"})
    base = await db.equipment.count_documents({"placement": "Base"})
    job = await db.equipment.count_documents({"placement": "Job"})
    open_mnt = await db.maintenance.count_documents({"status": "Open"})
    active_jobs = await db.jobs.count_documents({"status": {"$in": ["Active", "Open", "In Progress"]}})
    low_stock = await db.inventory_items.count_documents({"$expr": {"$lte": ["$stock", "$min_stock"]}})
    return ["Metric", "Value"], [
        ["Total Equipment", total], ["Green Tag / Ready", green],
        ["Red Tag / Under Maintenance", red], ["At Base", base], ["On Job", job],
        ["Open Maintenance", open_mnt], ["Active Jobs", active_jobs], ["Low Stock Items", low_stock],
    ]


async def _dataset(dataset, params, user):
    loaders = {
        "equipment": _equipment_rows,
        "maintenance": _maintenance_rows,
        "inventory": _inventory_rows,
        "calibration": _calibration_rows,
        "clients": _client_rows,
        "jobs": _job_rows,
        "audit": _audit_rows,
        "dashboard": _dashboard_rows,
    }
    if dataset == "users":
        return await _user_rows(params, user)
    loader = loaders.get(dataset)
    if not loader:
        raise HTTPException(status_code=404, detail="Export dataset not found")
    return await loader(params)


def _xlsx_bytes(
    title,
    headers,
    rows,
    timezone_name,
    brand_logo_bytes=None,
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.sheet_view.showGridLines = False

    generated = _local_now(
        timezone_name
    )

    data_column_count = max(
        1,
        len(headers),
    )
    layout_column_count = max(
        data_column_count,
        6,
    )
    layout_last = get_column_letter(
        layout_column_count
    )
    data_last = get_column_letter(
        data_column_count
    )

    ws.merge_cells(
        f"D1:{layout_last}1"
    )
    ws["D1"] = (
        f"AMT - {title} Export"
    )
    ws["D1"].font = Font(
        bold=True,
        size=18,
        color="0F172A",
    )
    ws["D1"].alignment = Alignment(
        vertical="center",
    )

    ws.merge_cells(
        f"D2:{layout_last}2"
    )
    ws["D2"] = (
        "Generated "
        f"{generated.strftime('%Y-%m-%d %H:%M %Z')} "
        f"- {len(rows)} record(s)"
    )
    ws["D2"].font = Font(
        size=9,
        color="64748B",
    )

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 8

    logo_source = None
    if brand_logo_bytes:
        logo_source = io.BytesIO(
            brand_logo_bytes
        )
    elif AMT_MARK_TAGLINE.exists():
        logo_source = str(
            AMT_MARK_TAGLINE
        )

    if logo_source is not None:
        try:
            logo = XLImage(
                logo_source
            )
            scale = min(
                180 / logo.width,
                62 / logo.height,
            )
            logo.width = int(
                logo.width * scale
            )
            logo.height = int(
                logo.height * scale
            )
            ws.add_image(
                logo,
                "A1",
            )
        except Exception:
            pass

    header_row = 5

    thin = Side(
        style="thin",
        color="CBD5E1",
    )
    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    for index, value in enumerate(
        headers,
        start=1,
    ):
        cell = ws.cell(
            row=header_row,
            column=index,
            value=_safe_excel(value),
        )
        cell.font = Font(
            bold=True,
            color="FFFFFF",
            size=9,
        )
        cell.fill = PatternFill(
            "solid",
            fgColor="0F172A",
        )
        cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border

    ws.row_dimensions[
        header_row
    ].height = 28

    for row_index, values in enumerate(
        rows,
        start=header_row + 1,
    ):
        for column_index, value in enumerate(
            values,
            start=1,
        ):
            cell = ws.cell(
                row=row_index,
                column=column_index,
                value=_safe_excel(value),
            )
            cell.font = Font(
                size=9,
                color="0F172A",
            )
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )
            cell.border = border

            if (
                row_index - header_row
            ) % 2 == 0:
                cell.fill = PatternFill(
                    "solid",
                    fgColor="F8FAFC",
                )

    for column_index in range(
        1,
        data_column_count + 1,
    ):
        letter = get_column_letter(
            column_index
        )
        max_len = len(
            str(
                headers[
                    column_index - 1
                ]
                or ""
            )
        )

        for row_index in range(
            header_row + 1,
            min(
                header_row
                + len(rows)
                + 1,
                header_row + 301,
            ),
        ):
            value = ws.cell(
                row=row_index,
                column=column_index,
            ).value
            max_len = max(
                max_len,
                len(
                    str(
                        value
                        or ""
                    )
                ),
            )

        ws.column_dimensions[
            letter
        ].width = min(
            max(
                max_len + 2,
                11,
            ),
            32,
        )

    if title.lower() == "maintenance":
        preferred = {
            5: 24,
            8: 24,
            12: 30,
            13: 30,
        }
        for column_index, width in preferred.items():
            if column_index <= data_column_count:
                ws.column_dimensions[
                    get_column_letter(
                        column_index
                    )
                ].width = width

    ws.freeze_panes = (
        f"A{header_row + 1}"
    )

    if rows:
        ws.auto_filter.ref = (
            f"A{header_row}:"
            f"{data_last}"
            f"{header_row + len(rows)}"
        )

    ws.print_title_rows = (
        f"1:{header_row}"
    )
    ws.print_area = (
        f"A1:{data_last}"
        f"{header_row + max(1, len(rows))}"
    )
    ws.page_setup.orientation = (
        "landscape"
    )
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.45
    ws.oddFooter.center.text = (
        "Powered by AMT "
        "(Asset Maintenance Tracker) "
        "- LogiSource Digital"
    )
    ws.oddFooter.right.text = (
        "Page &P of &N"
    )

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.read()


def _pdf_bytes(title, headers, rows, timezone_name, brand_logo_bytes=None):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=12 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ExportTitle",
        parent=styles["Heading1"],
        fontSize=15,
        leading=17,
        textColor=colors.HexColor("#0F172A"),
        spaceAfter=2,
    )

    sub_style = ParagraphStyle(
        "ExportSub",
        parent=styles["Normal"],
        fontSize=7,
        leading=9,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=0,
    )

    cell_style = ParagraphStyle(
        "ExportCell",
        parent=styles["Normal"],
        fontSize=6.2,
        leading=7.5,
        textColor=colors.HexColor("#0F172A"),
    )

    head_style = ParagraphStyle(
        "ExportHead",
        parent=cell_style,
        textColor=colors.white,
        fontName="Helvetica-Bold",
    )

    generated = _local_now(timezone_name)

    title_block = [
        Paragraph(
            f"AMT - {escape(title)} Export",
            title_style,
        ),
        Paragraph(
            (
                f"Generated "
                f"{generated.strftime('%Y-%m-%d %H:%M %Z')} "
                f"- {len(rows)} record(s)"
            ),
            sub_style,
        ),
    ]

    if brand_logo_bytes:
        try:
            logo = RLImage(
                io.BytesIO(
                    brand_logo_bytes
                )
            )
            max_w = 42 * mm
            max_h = 18 * mm
            scale = min(
                max_w / logo.imageWidth,
                max_h / logo.imageHeight,
            )
            logo.drawWidth = (
                logo.imageWidth
                * scale
            )
            logo.drawHeight = (
                logo.imageHeight
                * scale
            )
            logo.hAlign = "RIGHT"
            logo_cell = logo
        except Exception:
            brand_logo_bytes = None

    if (
        not brand_logo_bytes
        and AMT_MARK_TAGLINE.exists()
    ):
        logo_width = 40 * mm
        logo_height = logo_width * (
            835 / 1883
        )
        logo = RLImage(
            str(AMT_MARK_TAGLINE),
            width=logo_width,
            height=logo_height,
        )
        logo.hAlign = "RIGHT"
        logo_cell = logo
    elif not brand_logo_bytes:
        logo_cell = ""

    # One header row keeps the title/generated block and AMT logo
    # vertically aligned. The data table begins immediately below it.
    header_table = Table(
        [[title_block, logo_cell]],
        colWidths=[225 * mm, 52 * mm],
    )
    header_table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (1, 0), (1, 0), "RIGHT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]
        )
    )

    story = [
        header_table,
        Spacer(1, 4 * mm),
    ]

    data = [
        [
            Paragraph(
                _safe_pdf(value),
                head_style,
            )
            for value in headers
        ]
    ]

    data.extend(
        [
            [
                Paragraph(
                    _safe_pdf(value),
                    cell_style,
                )
                for value in row
            ]
            for row in rows
        ]
    )

    available_width = 277 * mm
    count = max(1, len(headers))
    widths = [
        available_width / count
    ] * count

    table = Table(
        data,
        colWidths=widths,
        repeatRows=1,
    )

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#0F172A"),
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.HexColor("#CBD5E1"),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor("#F8FAFC"),
                    ],
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    2,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    2,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    2,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    2,
                ),
            ]
        )
    )

    story.append(table)
    story.append(
        Spacer(
            1,
            4 * mm,
        )
    )
    story.append(
        Paragraph(
            (
                "Powered by AMT (Asset Maintenance Tracker) "
                "- LogiSource Digital"
            ),
            sub_style,
        )
    )
    doc.build(story)

    buf.seek(0)
    return buf.read()


def _params(**kwargs):
    return kwargs


@router.get("/export/{dataset}.xlsx")
async def export_xlsx(
    dataset: str, q: str = "", status: str = "", placement: str = "", type: str = "",
    category: str = "", low: str = "", entity_type: str = "", sap_no: str = "", serial_no: str = "",
    technician: str = "", failure: str = "", client_id: str = "", job_id: str = "",
    date_from: str = "", date_to: str = "", maintenance_ids: str = "",
    user: dict = Depends(get_current_user),
):
    await require_feature_enabled(
        "export_data"
    )

    params = _params(q=q, status=status, placement=placement, type=type, category=category, low=low,
                     entity_type=entity_type, sap_no=sap_no, serial_no=serial_no,
                     technician=technician, failure=failure, client_id=client_id,
                     job_id=job_id, date_from=date_from, date_to=date_to,
                     maintenance_ids=maintenance_ids)
    headers, rows = await _dataset(dataset, params, user)
    brand_logo = (
        await get_pdf_brand_logo_bytes()
    )
    data = _xlsx_bytes(
        dataset.title(),
        headers,
        rows,
        await _timezone_name(),
        brand_logo,
    )
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="amt-{dataset}-export.xlsx"'},
    )


@router.get("/export/{dataset}.pdf")
async def export_pdf(
    dataset: str, q: str = "", status: str = "", placement: str = "", type: str = "",
    category: str = "", low: str = "", entity_type: str = "", sap_no: str = "", serial_no: str = "",
    technician: str = "", failure: str = "", client_id: str = "", job_id: str = "",
    date_from: str = "", date_to: str = "", maintenance_ids: str = "",
    user: dict = Depends(get_current_user),
):
    await require_feature_enabled(
        "export_data"
    )

    params = _params(q=q, status=status, placement=placement, type=type, category=category, low=low,
                     entity_type=entity_type, sap_no=sap_no, serial_no=serial_no,
                     technician=technician, failure=failure, client_id=client_id,
                     job_id=job_id, date_from=date_from, date_to=date_to,
                     maintenance_ids=maintenance_ids)
    headers, rows = await _dataset(dataset, params, user)
    brand_logo = (
        await get_pdf_brand_logo_bytes()
    )
    data = _pdf_bytes(
        dataset.title(),
        headers,
        rows,
        await _timezone_name(),
        brand_logo,
    )
    return Response(
        content=data,
        media_type="application/pdf",
        headers={"Content-Disposition": f'inline; filename="amt-{dataset}-export.pdf"'},
    )
