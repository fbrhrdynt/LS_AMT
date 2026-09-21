import io
import re

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from calibration_routes import compute_expired_date
from core import audit_log, db, new_id, now_iso
from importer import to_iso_date


SUPPORTED_IMPORT_DATASETS = {
    "equipment",
    "calibration",
    "inventory",
}


DATASET_SPECS = {
    "equipment": {
        "sheet": "Equipment",
        "title": "AMT Equipment Import Template",
        "headers": [
            "SAP No.",
            "Mfg / Serial No.",
            "Equipment Name",
            "Category",
            "Manufacturer",
            "Date of Purchase",
            "Physical Condition",
        ],
        "example": [
            "14200988",
            "20237",
            "CENTRIFUGE BC20 RADIAL",
            "Centrifuge",
            "G-TECH",
            "2026-01-15",
            "Good",
        ],
        "required": ["SAP No."],
    },
    "calibration": {
        "sheet": "Calibration",
        "title": "AMT Calibration Import Template",
        "headers": [
            "Tool ID / Serial",
            "Tool Name",
            "Category",
            "Manufacturer",
            "Model",
            "Range / Set Pressure",
            "Calibration Date",
            "Frequency Value",
            "Frequency Unit",
            "Certificate No.",
            "Calibrated By",
            "Assigned SAP",
            "Equipment Name",
            "Comments / Notes",
        ],
        "example": [
            "CLPA21",
            "Pressure Gauge",
            "Pressure",
            "WIKA",
            "232.50",
            "0-10 bar",
            "2026-08-08",
            52,
            "week",
            "CERT-2026-001",
            "Approved Lab",
            "14200988",
            "CENTRIFUGE BC20 RADIAL",
            "Imported from calibration register",
        ],
        "required": [
            "Tool ID / Serial",
            "Tool Name",
            "Category",
        ],
    },
    "inventory": {
        "sheet": "Inventory",
        "title": "AMT Inventory Import Template",
        "headers": [
            "Item Code",
            "Item Name",
            "Category",
            "Type",
            "Part Number",
            "Unit",
            "Stock",
            "Minimum Stock",
            "Storage Location",
            "Unit Price",
        ],
        "example": [
            "SP-1001",
            "Bearing Set",
            "Centrifuge",
            "Spare Part",
            "BRG-BC20",
            "SET",
            10,
            3,
            "Rack A1",
            125.50,
        ],
        "required": [
            "Item Code",
            "Item Name",
        ],
    },
}


def _norm_header(value):
    return re.sub(
        r"[^a-z0-9]+",
        "",
        str(value or "").strip().lower(),
    )


HEADER_MAPS = {
    "equipment": {
        "sapno": "sap_no",
        "assetno": "sap_no",
        "mfgserialno": "mfg_no",
        "serialmfgno": "mfg_no",
        "serialno": "mfg_no",
        "equipmentname": "name",
        "equipment": "name",
        "category": "category",
        "manufacturer": "manufacturer",
        "dateofpurchase": "date_of_purchase",
        "physicalcondition": "physical_condition",
    },
    "calibration": {
        "toolidserial": "tool_id",
        "toolid": "tool_id",
        "serialno": "tool_id",
        "toolname": "tool_name",
        "category": "category",
        "manufacturer": "manufacturer",
        "model": "model",
        "rangesetpressure": "range_spec",
        "range": "range_spec",
        "calibrationdate": "calibration_date",
        "frequencyvalue": "frequency_value",
        "frequencyunit": "frequency_unit",
        "certificateno": "cert_number",
        "certnumber": "cert_number",
        "calibratedby": "calibrated_by",
        "assignedsap": "sap_no",
        "sapno": "sap_no",
        "equipmentname": "equipment_name",
        "commentsnotes": "comments",
        "comments": "comments",
        "notes": "comments",
    },
    "inventory": {
        "itemcode": "item_code",
        "itemname": "item_name",
        "category": "category",
        "type": "type",
        "partnumber": "part_number",
        "unit": "unit",
        "stock": "stock",
        "minimumstock": "min_stock",
        "minstock": "min_stock",
        "storagelocation": "storage_location",
        "location": "storage_location",
        "unitprice": "unit_price",
    },
}


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _float(value, default=0.0):
    if value in (None, ""):
        return float(default)
    try:
        return float(value)
    except (TypeError, ValueError):
        raise ValueError(f"must be a number, got {value!r}")


def _int(value, default=0):
    if value in (None, ""):
        return int(default)
    try:
        return int(float(value))
    except (TypeError, ValueError):
        raise ValueError(f"must be a whole number, got {value!r}")


def _template_notes(dataset):
    common = [
        "Keep the header names unchanged.",
        "Start entering real data on row 5 of the data sheet.",
        "Blank rows are ignored.",
        "Duplicate records can be skipped during the validation/import workflow.",
    ]
    specific = {
        "equipment": [
            "SAP No. is required and is matched against the existing Equipment register.",
            "New equipment is created at Base with Operational status.",
            "Date of Purchase should use YYYY-MM-DD where possible.",
        ],
        "calibration": [
            "Tool ID / Serial, Tool Name, and Category are required.",
            "Frequency Unit must be week or month.",
            "Assigned SAP is optional.",
            "If Assigned SAP does not exist, AMT will ask whether to create the Equipment during the next step.",
            "New Equipment created from Calibration import is registered at Base.",
            "Certificate files are still uploaded from the Calibration menu.",
        ],
        "inventory": [
            "Item Code and Item Name are required.",
            "Type should be Spare Part or Consumable.",
            "Stock, Minimum Stock, and Unit Price must be numeric.",
            "Initial stock creates an Inventory initial-stock transaction.",
        ],
    }
    return common + specific[dataset]


def build_import_template(dataset: str) -> bytes:
    dataset = str(dataset or "").strip().lower()
    spec = DATASET_SPECS.get(dataset)
    if not spec:
        raise ValueError("Unsupported import template")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = spec["sheet"]
    ws.sheet_view.showGridLines = False

    last_col = openpyxl.utils.get_column_letter(
        len(spec["headers"])
    )

    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = spec["title"]
    ws["A1"].font = Font(
        bold=True,
        size=16,
        color="0F172A",
    )
    ws["A1"].alignment = Alignment(
        vertical="center",
    )
    ws.row_dimensions[1].height = 28

    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = (
        "Fill from row 5, save as .xlsx, then upload through AMT > Excel Import."
    )
    ws["A2"].font = Font(
        size=9,
        color="64748B",
    )

    header_row = 4
    thin = Side(
        style="thin",
        color="CBD5E1",
    )
    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    for index, header in enumerate(
        spec["headers"],
        start=1,
    ):
        cell = ws.cell(
            row=header_row,
            column=index,
            value=header,
        )
        cell.font = Font(
            bold=True,
            color="FFFFFF",
            size=9,
        )
        cell.fill = PatternFill(
            "solid",
            fgColor="0F172A",
        )
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )
        cell.border = border

        letter = openpyxl.utils.get_column_letter(index)
        ws.column_dimensions[letter].width = min(
            max(
                len(header) + 3,
                len(str(spec["example"][index - 1])) + 3,
                12,
            ),
            28,
        )

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{last_col}4"

    if dataset == "inventory":
        type_col = spec["headers"].index("Type") + 1
        letter = openpyxl.utils.get_column_letter(type_col)
        dv = DataValidation(
            type="list",
            formula1='"Spare Part,Consumable"',
            allow_blank=True,
        )
        ws.add_data_validation(dv)
        dv.add(f"{letter}5:{letter}1000")

    if dataset == "calibration":
        unit_col = spec["headers"].index("Frequency Unit") + 1
        letter = openpyxl.utils.get_column_letter(unit_col)
        dv = DataValidation(
            type="list",
            formula1='"week,month"',
            allow_blank=False,
        )
        ws.add_data_validation(dv)
        dv.add(f"{letter}5:{letter}1000")

    info = wb.create_sheet("Instructions")
    info["A1"] = "AMT Import Instructions"
    info["A1"].font = Font(
        bold=True,
        size=16,
        color="0F172A",
    )
    info["A3"] = "Dataset"
    info["B3"] = dataset.title()
    info["A4"] = "Required Fields"
    info["B4"] = ", ".join(spec["required"])

    row = 6
    for note in _template_notes(dataset):
        info.cell(row=row, column=1, value="•")
        info.cell(row=row, column=2, value=note)
        row += 1

    row += 1
    info.cell(row=row, column=1, value="Example")
    info.cell(row=row, column=1).font = Font(
        bold=True,
        color="0F172A",
    )
    row += 1

    for col, header in enumerate(
        spec["headers"],
        start=1,
    ):
        cell = info.cell(
            row=row,
            column=col,
            value=header,
        )
        cell.font = Font(
            bold=True,
            color="FFFFFF",
            size=9,
        )
        cell.fill = PatternFill(
            "solid",
            fgColor="334155",
        )
        cell.alignment = Alignment(
            wrap_text=True,
        )

    row += 1
    for col, value in enumerate(
        spec["example"],
        start=1,
    ):
        info.cell(
            row=row,
            column=col,
            value=value,
        ).alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

    info.column_dimensions["A"].width = 18
    for col in range(2, len(spec["headers"]) + 1):
        info.column_dimensions[
            openpyxl.utils.get_column_letter(col)
        ].width = 18

    for r in range(6, row + 1):
        info.cell(row=r, column=2).alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.read()


def _find_header_row(ws, dataset):
    mapping = HEADER_MAPS[dataset]
    required_norm = {
        _norm_header(value)
        for value in DATASET_SPECS[dataset]["required"]
    }

    for row_index in range(
        1,
        min(ws.max_row, 12) + 1,
    ):
        values = [
            _norm_header(
                ws.cell(
                    row=row_index,
                    column=col,
                ).value
            )
            for col in range(
                1,
                ws.max_column + 1,
            )
        ]
        recognized = {
            value
            for value in values
            if value in mapping
        }
        if required_norm.intersection(recognized):
            return row_index

    raise ValueError(
        "Could not find the AMT template header row. "
        "Download a fresh template and keep its header names unchanged."
    )


def parse_import_template(
    path_or_bytes,
    dataset: str,
):
    dataset = str(dataset or "").strip().lower()
    if dataset not in SUPPORTED_IMPORT_DATASETS:
        raise ValueError("Unsupported import dataset")

    wb = openpyxl.load_workbook(
        path_or_bytes,
        data_only=True,
        read_only=True,
    )

    wanted = DATASET_SPECS[dataset]["sheet"].lower()
    sheet_name = next(
        (
            name
            for name in wb.sheetnames
            if name.lower() == wanted
        ),
        wb.sheetnames[0],
    )
    ws = wb[sheet_name]
    header_row = _find_header_row(
        ws,
        dataset,
    )

    mapping = HEADER_MAPS[dataset]
    columns = {}

    for col in range(
        1,
        ws.max_column + 1,
    ):
        key = _norm_header(
            ws.cell(
                row=header_row,
                column=col,
            ).value
        )
        canonical = mapping.get(key)
        if canonical:
            columns[col] = canonical

    rows = []
    errors = []

    for row_number in range(
        header_row + 1,
        ws.max_row + 1,
    ):
        raw = {
            canonical: ws.cell(
                row=row_number,
                column=col,
            ).value
            for col, canonical in columns.items()
        }

        if not any(
            value not in (None, "")
            for value in raw.values()
        ):
            continue

        try:
            if dataset == "equipment":
                row = {
                    "sap_no": _text(raw.get("sap_no")),
                    "mfg_no": _text(raw.get("mfg_no")),
                    "name": _text(raw.get("name")),
                    "category": _text(raw.get("category")),
                    "manufacturer": _text(raw.get("manufacturer")),
                    "date_of_purchase": to_iso_date(
                        raw.get("date_of_purchase")
                    ),
                    "physical_condition": _text(
                        raw.get("physical_condition")
                    ),
                }
                if not row["sap_no"]:
                    raise ValueError("SAP No. is required")

            elif dataset == "calibration":
                frequency_unit = (
                    _text(
                        raw.get("frequency_unit")
                    )
                    or "week"
                ).lower()
                if frequency_unit in ("weeks", "weekly"):
                    frequency_unit = "week"
                if frequency_unit in ("months", "monthly"):
                    frequency_unit = "month"
                if frequency_unit not in ("week", "month"):
                    raise ValueError(
                        "Frequency Unit must be week or month"
                    )

                row = {
                    "tool_id": _text(raw.get("tool_id")),
                    "tool_name": _text(raw.get("tool_name")),
                    "category": _text(raw.get("category")),
                    "manufacturer": _text(raw.get("manufacturer")),
                    "model": _text(raw.get("model")),
                    "range_spec": _text(raw.get("range_spec")),
                    "calibration_date": (
                        to_iso_date(
                            raw.get("calibration_date")
                        )
                        or ""
                    ),
                    "frequency_value": _int(
                        raw.get("frequency_value"),
                        52,
                    ),
                    "frequency_unit": frequency_unit,
                    "cert_number": _text(
                        raw.get("cert_number")
                    ),
                    "calibrated_by": _text(
                        raw.get("calibrated_by")
                    ),
                    "sap_no": _text(raw.get("sap_no")),
                    "equipment_name": _text(
                        raw.get("equipment_name")
                    ),
                    "comments": _text(raw.get("comments")),
                }
                if not row["tool_id"]:
                    raise ValueError(
                        "Tool ID / Serial is required"
                    )
                if not row["tool_name"]:
                    raise ValueError(
                        "Tool Name is required"
                    )
                if not row["category"]:
                    raise ValueError(
                        "Category is required"
                    )
                if not 1 <= row["frequency_value"] <= 520:
                    raise ValueError(
                        "Frequency Value must be between 1 and 520"
                    )

            else:
                item_type = (
                    _text(raw.get("type"))
                    or "Spare Part"
                )
                if item_type.lower() in (
                    "spare part",
                    "sparepart",
                    "spare",
                ):
                    item_type = "Spare Part"
                elif item_type.lower() in (
                    "consumable",
                    "consumables",
                ):
                    item_type = "Consumable"
                else:
                    raise ValueError(
                        "Type must be Spare Part or Consumable"
                    )

                row = {
                    "item_code": _text(
                        raw.get("item_code")
                    ),
                    "item_name": _text(
                        raw.get("item_name")
                    ),
                    "category": _text(
                        raw.get("category")
                    ),
                    "type": item_type,
                    "part_number": _text(
                        raw.get("part_number")
                    ),
                    "unit": (
                        _text(raw.get("unit"))
                        or "EA"
                    ),
                    "stock": _float(
                        raw.get("stock"),
                        0,
                    ),
                    "min_stock": _float(
                        raw.get("min_stock"),
                        0,
                    ),
                    "storage_location": _text(
                        raw.get("storage_location")
                    ),
                    "unit_price": _float(
                        raw.get("unit_price"),
                        0,
                    ),
                }
                if not row["item_code"]:
                    raise ValueError(
                        "Item Code is required"
                    )
                if not row["item_name"]:
                    raise ValueError(
                        "Item Name is required"
                    )

            row["_row_number"] = row_number
            rows.append(row)

        except ValueError as exc:
            errors.append(
                {
                    "row": row_number,
                    "message": str(exc),
                }
            )

    wb.close()
    return {
        "dataset": dataset,
        "rows": rows,
        "errors": errors,
    }


def _public_row(row):
    return {
        key: value
        for key, value in row.items()
        if not key.startswith("_")
    }


async def analyze_import_rows(parsed):
    dataset = parsed["dataset"]
    rows = parsed["rows"]
    errors = parsed["errors"]

    duplicates = 0
    missing_equipment = []

    if dataset == "equipment":
        existing = set(
            value
            for value in await db.equipment.distinct(
                "sap_no"
            )
            if value
        )
        duplicates = sum(
            1
            for row in rows
            if row["sap_no"] in existing
        )

    elif dataset == "inventory":
        existing = set(
            value
            for value in await db.inventory_items.distinct(
                "item_code"
            )
            if value
        )
        duplicates = sum(
            1
            for row in rows
            if row["item_code"] in existing
        )

    else:
        current = await db.calibration_tools.find(
            {
                "is_deleted": {"$ne": True},
            },
            {
                "_id": 0,
                "tool_id": 1,
                "category": 1,
            },
        ).to_list(100000)
        existing = {
            (
                _text(item.get("tool_id")).casefold(),
                _text(item.get("category")).casefold(),
            )
            for item in current
        }
        duplicates = sum(
            1
            for row in rows
            if (
                row["tool_id"].casefold(),
                row["category"].casefold(),
            )
            in existing
        )

        existing_saps = set(
            value
            for value in await db.equipment.distinct(
                "sap_no"
            )
            if value
        )
        missing = {}
        for row in rows:
            sap = row.get("sap_no")
            if sap and sap not in existing_saps:
                missing.setdefault(
                    sap,
                    {
                        "sap_no": sap,
                        "equipment_name": (
                            row.get("equipment_name")
                            or f"Imported Equipment {sap}"
                        ),
                    },
                )
        missing_equipment = list(
            missing.values()
        )

    return {
        "dataset": dataset,
        "total": len(rows) + len(errors),
        "valid": len(rows),
        "invalid": len(errors),
        "new": max(
            0,
            len(rows) - duplicates,
        ),
        "duplicates": duplicates,
        "sample": [
            _public_row(row)
            for row in rows[:8]
        ],
        "errors": errors[:50],
        "missing_equipment": missing_equipment,
    }


async def _create_equipment(
    row,
    user_name,
    source="template-import",
):
    sap = _text(row.get("sap_no"))
    if not sap:
        raise ValueError("SAP No. is required")

    doc = {
        "id": new_id(),
        "sap_no": sap,
        "mfg_no": _text(row.get("mfg_no")),
        "name": (
            _text(row.get("name"))
            or _text(row.get("equipment_name"))
            or f"Imported Equipment {sap}"
        ),
        "category": _text(row.get("category")),
        "manufacturer": _text(row.get("manufacturer")),
        "date_of_purchase": row.get("date_of_purchase"),
        "physical_condition": _text(
            row.get("physical_condition")
        ),
        "placement": "Base",
        "placement_detail": "Base",
        "operational_status": "Operational",
        "current_job_id": None,
        "current_client_id": None,
        "source": source,
        "created_at": now_iso(),
        "updated_at": now_iso(),
    }

    await db.equipment.insert_one(doc)
    await db.location_history.insert_one(
        {
            "id": new_id(),
            "equipment_id": doc["id"],
            "from_placement": None,
            "to_placement": "Base",
            "placement_detail": "Base",
            "job_id": None,
            "reason": "Initial registration from Excel import",
            "created_by": user_name,
            "created_at": now_iso(),
        }
    )
    return doc


async def _create_inventory(
    row,
    user_name,
):
    now = now_iso()
    doc = {
        "id": new_id(),
        "item_code": row["item_code"],
        "item_name": row["item_name"],
        "category": row.get("category", ""),
        "type": row.get("type", "Spare Part"),
        "part_number": row.get("part_number", ""),
        "unit": row.get("unit", "EA"),
        "stock": float(row.get("stock") or 0),
        "min_stock": float(row.get("min_stock") or 0),
        "storage_location": row.get(
            "storage_location",
            "",
        ),
        "unit_price": float(
            row.get("unit_price") or 0
        ),
        "created_at": now,
        "updated_at": now,
    }
    await db.inventory_items.insert_one(doc)

    if doc["stock"]:
        await db.inventory_transactions.insert_one(
            {
                "id": new_id(),
                "item_id": doc["id"],
                "item_code": doc["item_code"],
                "item_name": doc["item_name"],
                "type": "initial",
                "direction": "in",
                "qty": doc["stock"],
                "unit": doc["unit"],
                "maintenance_id": None,
                "equipment_id": None,
                "balance_after": doc["stock"],
                "note": "Initial stock from Excel import",
                "created_by": user_name,
                "created_at": now,
            }
        )
    return doc


async def _create_calibration(
    row,
    equipment,
    user_name,
):
    now = now_iso()
    expired_date = compute_expired_date(
        row.get("calibration_date") or "",
        int(row.get("frequency_value") or 52),
        row.get("frequency_unit") or "week",
    )

    doc = {
        "id": new_id(),
        "tool_id": row["tool_id"],
        "tool_name": row["tool_name"],
        "category": row["category"],
        "manufacturer": row.get("manufacturer", ""),
        "model": row.get("model", ""),
        "range_spec": row.get("range_spec", ""),
        "calibration_date": row.get(
            "calibration_date",
            "",
        ),
        "frequency_value": int(
            row.get("frequency_value") or 52
        ),
        "frequency_unit": row.get(
            "frequency_unit",
            "week",
        ),
        "expired_date": expired_date,
        "cert_number": row.get("cert_number", ""),
        "calibrated_by": row.get("calibrated_by", ""),
        "comments": row.get("comments", ""),
        "equipment_id": (
            equipment.get("id")
            if equipment
            else None
        ),
        "equipment_sap_no": (
            equipment.get("sap_no")
            if equipment
            else None
        ),
        "equipment_name": (
            equipment.get("name")
            if equipment
            else None
        ),
        "is_deleted": False,
        "source": "template-import",
        "created_by": user_name,
        "created_at": now,
        "updated_at": now,
    }

    await db.calibration_tools.insert_one(doc)

    if equipment:
        await db.calibration_assignments.insert_one(
            {
                "id": new_id(),
                "calibration_tool_id": doc["id"],
                "tool_id": doc["tool_id"],
                "equipment_id": equipment["id"],
                "equipment_sap_no": equipment.get("sap_no"),
                "equipment_name": equipment.get("name"),
                "status": "Active",
                "assigned_at": now,
                "assigned_by": user_name,
            }
        )

    return doc


async def execute_import_rows(
    parsed,
    user,
    skip_duplicates=True,
    create_missing_saps=None,
):
    dataset = parsed["dataset"]
    rows = parsed["rows"]
    errors = parsed["errors"]

    if errors:
        raise ValueError(
            f"Workbook contains {len(errors)} invalid row(s). "
            "Fix the workbook and validate again before importing."
        )

    create_missing_saps = set(
        create_missing_saps or []
    )

    result = {
        "dataset": dataset,
        "added": 0,
        "skipped_duplicates": 0,
        "skipped_missing_equipment": 0,
        "equipment_created": 0,
    }

    user_name = user.get("name") or "system"

    if dataset == "equipment":
        existing = set(
            value
            for value in await db.equipment.distinct(
                "sap_no"
            )
            if value
        )

        for row in rows:
            sap = row["sap_no"]
            if sap in existing:
                if skip_duplicates:
                    result["skipped_duplicates"] += 1
                    continue
                raise ValueError(
                    f"Duplicate Equipment SAP No. {sap}"
                )

            await _create_equipment(
                row,
                user_name,
            )
            existing.add(sap)
            result["added"] += 1

    elif dataset == "inventory":
        existing = set(
            value
            for value in await db.inventory_items.distinct(
                "item_code"
            )
            if value
        )

        for row in rows:
            code = row["item_code"]
            if code in existing:
                if skip_duplicates:
                    result["skipped_duplicates"] += 1
                    continue
                raise ValueError(
                    f"Duplicate Inventory Item Code {code}"
                )

            await _create_inventory(
                row,
                user_name,
            )
            existing.add(code)
            result["added"] += 1

    else:
        equipment_rows = await db.equipment.find(
            {},
            {"_id": 0},
        ).to_list(100000)
        equipment_map = {
            item["sap_no"]: item
            for item in equipment_rows
            if item.get("sap_no")
        }

        existing_cal = await db.calibration_tools.find(
            {
                "is_deleted": {"$ne": True},
            },
            {
                "_id": 0,
                "tool_id": 1,
                "category": 1,
            },
        ).to_list(100000)
        existing_keys = {
            (
                _text(item.get("tool_id")).casefold(),
                _text(item.get("category")).casefold(),
            )
            for item in existing_cal
        }

        for row in rows:
            key = (
                row["tool_id"].casefold(),
                row["category"].casefold(),
            )
            if key in existing_keys:
                if skip_duplicates:
                    result["skipped_duplicates"] += 1
                    continue
                raise ValueError(
                    "Duplicate Calibration Tool "
                    f"{row['tool_id']} / {row['category']}"
                )

            equipment = None
            sap = row.get("sap_no")

            if sap:
                equipment = equipment_map.get(sap)

                if (
                    not equipment
                    and sap in create_missing_saps
                ):
                    equipment = await _create_equipment(
                        {
                            "sap_no": sap,
                            "name": (
                                row.get("equipment_name")
                                or f"Imported Equipment {sap}"
                            ),
                        },
                        user_name,
                        source="calibration-template-import",
                    )
                    equipment_map[sap] = equipment
                    result["equipment_created"] += 1

                if not equipment:
                    result[
                        "skipped_missing_equipment"
                    ] += 1
                    continue

            await _create_calibration(
                row,
                equipment,
                user_name,
            )
            existing_keys.add(key)
            result["added"] += 1

    await audit_log(
        "import",
        dataset,
        "import.template.execute",
        user,
        (
            f"Imported {result['added']} {dataset} record(s); "
            f"duplicates skipped={result['skipped_duplicates']}; "
            f"missing equipment skipped={result['skipped_missing_equipment']}; "
            f"equipment created={result['equipment_created']}"
        ),
    )

    return result
